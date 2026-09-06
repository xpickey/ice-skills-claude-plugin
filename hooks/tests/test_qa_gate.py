#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""ทดสอบ ice-qa-round-gate.py — รอบ 1–3 ผ่าน · รอบ 4 ปฏิเสธ · ICE_QA_EXTEND=1 ผ่านหนึ่งรอบ · ไฟล์คนละชิ้นนับแยก · agent อื่นไม่ถูกแตะ"""
import json
import os
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
HOOKS = os.path.dirname(HERE)
sys.path.insert(0, HOOKS)
import ice_route_lib as lib  # noqa: E402
GATE = os.path.join(HOOKS, "ice-qa-round-gate.py")


def call(sid, agent, prompt):
    p = subprocess.run([sys.executable, GATE], input=json.dumps({"session_id": sid, "tool_name": "Agent", "tool_input": {"subagent_type": agent, "prompt": prompt}}), capture_output=True, text=True)
    return '"deny"' in p.stdout


def case_final(name, prompt, expect_deny):
    """ด่าน is_final: งานฉบับสุดท้ายต้องไม่มีข้อความรอเติม (ซ้อมจริง 2026.09.06)"""
    p = subprocess.run([sys.executable, GATE], input=json.dumps({"session_id": "finaltest", "tool_input": {"subagent_type": "qa-master-agent", "prompt": prompt}}), capture_output=True, text=True)
    denied = "deny" in p.stdout
    ok = denied == expect_deny
    print(("  ✓ " if ok else "  ✗ ") + name + ("" if ok else f"  (deny={denied} expected={expect_deny})"))
    return 0 if ok else 1


def main():
    sid = "qatest-" + next(tempfile._get_candidate_names())
    bad = 0
    deck = "ตรวจ /x/Deck_V01R03_2026.09.05.pptx ตาม spec"
    for i in range(1, 4):
        d = call(sid, "qa-master-agent", deck.replace("R03", f"R0{i}"))
        ok = not d; bad += 0 if ok else 1; print(("  ✓ " if ok else "  ✗ ") + f"รอบ {i} ผ่าน")
    d = call(sid, "qa-master-agent", deck.replace("R03", "R04")); ok = d; bad += 0 if ok else 1; print(("  ✓ " if ok else "  ✗ ") + "รอบ 4 ถูกปฏิเสธ")
    d = call(sid, "qa-master-agent", "ICE_QA_EXTEND=1 " + deck.replace("R03", "R04")); ok = not d; bad += 0 if ok else 1; print(("  ✓ " if ok else "  ✗ ") + "user อนุมัติขยาย (ICE_QA_EXTEND=1) ผ่านหนึ่งรอบ")
    d = call(sid, "qa-master-agent", deck.replace("R03", "R05")); ok = d; bad += 0 if ok else 1; print(("  ✓ " if ok else "  ✗ ") + "รอบถัดไปโดยไม่มีอนุมัติ ถูกปฏิเสธอีก")
    d = call(sid, "qa-master-agent", "ตรวจ /x/Workbook_V01R01.xlsx"); ok = not d; bad += 0 if ok else 1; print(("  ✓ " if ok else "  ✗ ") + "ไฟล์คนละชิ้น นับแยก")
    d = call(sid, "solution-knowledge-agent", deck); ok = not d; bad += 0 if ok else 1; print(("  ✓ " if ok else "  ✗ ") + "agent อื่นไม่ถูกแตะ")
    # ด่าน is_final (ซ้อมจริง 2026.09.06) — ต้องอ่าน path เต็มได้ ARTIFACT ตัวเดิมจับเฉพาะชื่อไฟล์
    import tempfile as _tf
    from pptx import Presentation
    tmpd = _tf.mkdtemp(prefix="qafinal-")
    for nm, txt in (("WithHolder_V01R01.pptx", "รายได้ [รอตัวเลขจริงจากลูกค้า]"), ("Clean_V01R01.pptx", "ปิดงบเร็วขึ้นด้วยระบบงาน ERP เดียวกันทั้งกลุ่ม")):
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = txt
        prs.save(os.path.join(tmpd, nm))
    bad += case_final("is_final ที่ยังมีข้อความรอเติม → ปฏิเสธ", f"ตรวจ is_final: true artifact_path: {tmpd}/WithHolder_V01R01.pptx", True)
    bad += case_final("is_final ที่สะอาดแล้ว → ผ่าน", f"ตรวจ is_final: true artifact_path: {tmpd}/Clean_V01R01.pptx", False)
    bad += case_final("ไม่ใช่ is_final แม้มีข้อความรอเติม → ผ่าน", f"ตรวจ is_final: false artifact_path: {tmpd}/WithHolder_V01R01.pptx", False)
    try:
        os.remove(lib.state_path(sid))
        os.remove(lib.state_path("finaltest"))
    except OSError:
        pass
    print("\nผล: " + ("ผ่านทั้งหมด" if bad == 0 else f"ไม่ผ่าน {bad} ข้อ"))
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
